import base64
import io
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main as parser
import parser_core
from openpyxl import load_workbook
from reportlab.pdfgen import canvas


def pdf_bytes():
    out = io.BytesIO()
    c = canvas.Canvas(out)
    c.drawString(20, 700, 'Synthetic PDF fixture')
    c.save()
    return out.getvalue()


def rows():
    result = [parser.PensionRow('01.09.2026', 10000 + i, '09.2026', f'Company {i}', f'{i:012}', None)
              for i in range(1, 5)]
    result.append(parser.PensionRow('01.09.2026', 999999, '09.2026', 'Company 1', '000000000001', 'ВОЗВРАЩЕН'))
    return result


class IntegrationTests(unittest.TestCase):
    def test_pension_only_and_missing_files(self):
        pension = dict(iin='000000000002', fio='Test pension', rows=rows())
        with patch.object(parser, 'extract_any_credit_report') as credit, \
             patch.object(parser, 'parse_pension_pdf_tables', return_value=pension):
            result = parser_core.parse_report(pension_pdf=pdf_bytes())
            credit.assert_not_called()
        self.assertTrue(result['ok'])
        self.assertFalse(result['has_credit'])
        self.assertTrue(result['has_pension'])
        self.assertEqual(result['iin'], pension['iin'])
        self.assertGreater(result['srzp'], 0)
        self.assertEqual(result['active'], [])
        wb = load_workbook(io.BytesIO(base64.b64decode(result['excel_b64'])))
        self.assertIn('Company 4', wb['Анкета']['A46'].value)
        with self.assertRaises(ValueError):
            parser_core.parse_report()
        from fastapi.testclient import TestClient
        import service
        with patch.object(service.lic, 'check_token', return_value={'ok': True}), \
             patch.object(parser, 'parse_pension_pdf_tables', return_value=dict(iin='000000000002', rows=rows())):
            client = TestClient(service.app)
            self.assertEqual(client.post('/parse', data={'token': 'test'}).status_code, 400)
            self.assertEqual(client.post('/parse', data={'token': 'test'}, files={'pension': ('p.pdf', pdf_bytes(), 'application/pdf')}).status_code, 200)
            self.assertEqual(client.post('/parse', data={'token': 'test'}, files={'pension': ('p.pdf', b'', 'application/pdf')}).status_code, 400)

    def test_orgs_returns_selection(self):
        rs = rows()
        self.assertEqual(len(parser.list_pension_orgs(rs)), 4)
        selected = parser.select_pension_orgs(rs, '1,000000000004')
        self.assertEqual(len(selected), 2)
        self.assertEqual(parser.calc_total_srzp(rs), sum(parser.calc_pension_avg_vals([r.amount]) for r in rs[:4]))
        with self.assertRaises(ValueError):
            parser.select_pension_orgs(rs, '99')

    def test_contract_split_and_empty_cessionary(self):
        self.assertEqual(len(parser.split_contract_blocks('КОНТРАКТ 1ЗАЙМ\nКОНТРАКТ 2КРЕДИТ')), 2)
        self.assertIsNone(parser.parse_cessionary('Наименование цессионария :\nАО Чужой банк'))
        self.assertIn('Эксперт', parser.parse_cessionary('Наименование цессионария : ТОО "Эксперт плюс"'))

    def test_split_creditor_label_in_multicolumn_report(self):
        block = 'Источник ТОО "Микрофинансовая организация Общая сумма кредита: 145000 KZT\nинформации "Sofi Finance" (Софи финанс) Сумма просроченных взносов: 0 KZT\n(Кредитор)\nИнформация по состоянию на: 28.02.2023'
        name = parser.parse_creditor(block)
        self.assertIn('Sofi Finance', name)
        self.assertNotIn('145000', name)
        self.assertNotIn('Сумма', name)
        self.assertIsNone(parser.parse_creditor('Источник информации: —\nДата: 28.02.2023'))

    def test_wrapped_pdf_cleanup_success_and_error(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / 'signed.pdf'
            original = b'CMS-prefix' + pdf_bytes() + b'CMS-suffix'
            source.write_bytes(original)
            paths = []
            @parser.with_unwrapped_pdf
            def reader(path, fail=False):
                paths.append(path)
                self.assertTrue(path.read_bytes().startswith(b'%PDF-'))
                if fail:
                    raise ValueError('fixture error')
                return parser.extract_pages_text_pdfplumber(path)
            self.assertIn('Synthetic', reader(source)[0])
            with self.assertRaises(ValueError):
                reader(source, fail=True)
            self.assertTrue(all(not p.exists() for p in paths))
            self.assertEqual(source.read_bytes(), original)

    def test_web_adapter_workbook_and_selection(self):
        cd = dict(source='pkb', iin='000000000001', fio='Test', pkr=700,
                  active=[], completed_recent=[], completed_old=[], revoked=[])
        pension = dict(iin=cd['iin'], fio='Test', rows=rows())
        with patch.object(parser, 'extract_any_credit_report', return_value=cd), \
             patch.object(parser, 'parse_pension_pdf_tables', return_value=pension):
            result = parser_core.parse_report(pdf_bytes(), pdf_bytes())
        self.assertTrue(result['ok'])
        wb = load_workbook(io.BytesIO(base64.b64decode(result['excel_b64'])))
        self.assertEqual(wb.sheetnames, ['Анкета', 'Кредиты', 'Анализ'])
        self.assertIn('Company 4', wb['Анкета']['A46'].value)
        self.assertTrue(all(o['selected'] for o in result['pension_organizations']))
        pension = dict(iin=cd['iin'], fio='Test', rows=rows())
        with patch.object(parser, 'extract_any_credit_report', return_value=cd), \
             patch.object(parser, 'parse_pension_pdf_tables', return_value=pension):
            result = parser_core.parse_report(pdf_bytes(), pdf_bytes(), pension_orgs='4')
        self.assertEqual(result['srzp'], parser.calc_pension_avg_vals([10004]))
        self.assertEqual(sum(o['selected'] for o in result['pension_organizations']), 1)

    def test_pension_errors_are_visible(self):
        cd = dict(source='pkb', iin='000000000001', pkr=700, active=[])
        with patch.object(parser, 'extract_any_credit_report', return_value=cd), \
             patch.object(parser, 'parse_pension_pdf_tables', side_effect=ValueError('invalid')):
            self.assertFalse(parser_core.parse_report(pdf_bytes(), pdf_bytes())['ok'])
        with patch.object(parser, 'extract_any_credit_report', return_value=cd), \
             patch.object(parser, 'parse_pension_pdf_tables', return_value=dict(iin='000000000002',rows=rows())):
            result = parser_core.parse_report(pdf_bytes(), pdf_bytes())
            self.assertTrue(result["ok"])
            self.assertEqual(result["iin"], cd["iin"])
            self.assertGreater(result["srzp"], 0)

    def test_selected_credits_and_no_income_export(self):
        first = parser.parse_contract_block('КОНТРАКТ 1ЗАЙМ\nСумма периодического платежа: 10000 KZT')
        second = parser.parse_contract_block('КОНТРАКТ 2ЗАЙМ\nСумма периодического платежа: 25000 KZT')
        cd = dict(source='pkb', iin='000000000001', pkr=700, active=[first,second])
        with patch.object(parser, 'extract_any_credit_report', return_value=cd), \
             patch.object(parser, 'parse_pension_pdf_tables', return_value=dict(rows=rows())):
            r = parser_core.parse_report(pdf_bytes(),pdf_bytes(),pension_orgs='none',credit_indices='1')
        self.assertEqual(r['srzp'],0)
        self.assertEqual(r['monthly_load'],25000)
        self.assertEqual(len(r['active']),1)
        wb=load_workbook(io.BytesIO(base64.b64decode(r['excel_b64'])))
        self.assertEqual(str(wb['Кредиты']['E6'].value),'25000')
        with patch.object(parser, 'extract_any_credit_report', return_value=cd):
            r=parser_core.parse_report(pdf_bytes(),credit_indices='none')
            self.assertEqual(r['monthly_load'],0)
            self.assertEqual(r['active'],[])
            with self.assertRaises(ValueError):
                parser_core.parse_report(pdf_bytes(),credit_indices='99')

    def test_actual_credit_and_signed_pension_pdf(self):
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import Table, TableStyle
        pdfmetrics.registerFont(TTFont('FixtureFont', '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'))
        out = io.BytesIO()
        c = canvas.Canvas(out)
        c.setFont('FixtureFont', 9)
        lines = ['ПЕРСОНАЛЬНЫЙ КРЕДИТНЫЙ ОТЧЕТ', 'ИИН: 000000000001', 'ПКР: 700',
                 'ДЕЙСТВУЮЩИЕ ДОГОВОРА', 'КОНТРАКТ 1ЗАЙМ', 'Источник информации (Кредитор): АО Тест Банк',
                 'Сумма периодического платежа: 10000 KZT', 'Непогашенная сумма по кредиту: 50000 KZT',
                 'ЗАВЕРШЕННЫЕ ДОГОВОРЫ', 'ДОГОВОРЫ, ЗАВЕРШЕННЫЕ БОЛЕЕ 5 ЛЕТ НАЗАД', 'ОТОЗВАННЫЕ ДОГОВОРЫ']
        for i, line in enumerate(lines):
            c.drawString(20, 800 - i * 18, line)
        c.save()
        credit = out.getvalue()
        out = io.BytesIO()
        c = canvas.Canvas(out, pagesize=(900,600))
        c.setFont('FixtureFont',9)
        c.drawString(20,560,'ТАӘ/ФИО: Тест Клиент')
        c.drawString(20,540,'ЖСН/ИИН: 000000000001')
        data = [['Дата','','','','','Отправитель','БИН','','Сумма','Статус','Период']]
        data += [['01.09.2026','','','','',f'Орг {i}',f'{i:012}','','10000','ПРИНЯТ','09.2026'] for i in range(1,5)]
        data += [['01.09.2026','','','','','Орг 1','000000000001','','999999','ВОЗВРАЩЕН','09.2026']]
        table = Table(data,colWidths=[70,20,20,20,20,110,100,20,60,100,70])
        table.setStyle(TableStyle([('FONTNAME',(0,0),(-1,-1),'FixtureFont'),('FONTSIZE',(0,0),(-1,-1),8),('GRID',(0,0),(-1,-1),0.5,'black')]))
        table.wrapOn(c,850,450)
        table.drawOn(c,20,350)
        c.save()
        result = parser_core.parse_report(credit, b'CMS' + out.getvalue() + b'END', pension_orgs='4')
        self.assertTrue(result['ok'])
        self.assertEqual(result['monthly_load'],10000)
        self.assertEqual(result['srzp'],parser.calc_pension_avg_vals([10000]))
        self.assertEqual(len(result['pension_organizations']),4)

    def test_http_form_selection_and_validation(self):
        from fastapi.testclient import TestClient
        import service
        client = TestClient(service.app)
        with patch.object(service.lic, 'check_token', return_value={'ok': True}), \
             patch.object(parser_core, 'parse_report', return_value={'ok': True}) as run:
            r = client.post('/api/index/parse',data={'token':'fixture','pension_orgs':'1,4'},files={'file':('test.pdf',pdf_bytes(),'application/pdf')})
            self.assertEqual(r.status_code,200)
            self.assertEqual(run.call_args.kwargs['pension_orgs'],'1,4')
        with patch.object(service.lic, 'check_token', return_value={'ok': True}), \
             patch.object(parser_core, 'parse_report', side_effect=ValueError('Unknown organisation')):
            r = client.post('/parse',files={'file':('test.pdf',pdf_bytes(),'application/pdf')})
            self.assertEqual(r.status_code,422)
            self.assertEqual(r.json()['error'],'Unknown organisation')
        with patch.object(service.lic, 'check_token', return_value={'ok': False}):
            self.assertEqual(client.post('/parse',files={'file':('test.pdf',pdf_bytes())}).status_code,401)

if __name__ == '__main__':
    unittest.main()
