"""Printable A4 export model shared by the parser and offline QA."""
import re
from pathlib import Path
import main as parser


def client_filename(fio):
    name = re.sub(r'[\\/:*?"<>|\x00-\x1f]', ' ', str(fio or 'Анкета клиента'))
    name = re.sub(r'\s+', ' ', name).strip(' .')[:140]
    return (name or 'Анкета клиента') + '.xlsx'


class Sheet:
    def __init__(self, name):
        self.model = dict(name=name, rows=[], merges=[], breaks=[])
        self.height = 0

    def row(self, values, kind='body', height=23, spans=None):
        if self.height + height > 690 and self.model['rows']:
            self.page()
        n = len(self.model['rows']) + 1
        self.model['rows'].append(dict(values=values,kind=kind,height=height))
        for start,end in spans or []:
            self.model['merges'].append([n,start,end])
        self.height += height
        return n

    def full(self,text,kind='body',height=26):
        height=max(height,14*((len(str(text))+89)//90)+8)
        return self.row([text],kind,height,[(1,8)])

    def kv(self,label,value,kind='body',height=26):
        height=max(height,14*((len(str(value))+49)//50)+8,14*((len(str(label))+31)//32)+8)
        return self.row([label,None,None,value],kind,height,[(1,3),(4,8)])

    def section(self,title,required=100):
        if self.height + required > 690:
            self.page()
        self.full(title,'section',26)

    def page(self):
        if self.model['rows'] and len(self.model['rows']) not in self.model['breaks']:
            self.model['breaks'].append(len(self.model['rows']))
        self.height=0


def print_model(cd,pension,active,recent,old,revoked,srzp,has_credit,has_pension):
    fio=(pension or {}).get('fio') or cd.get('fio') or 'ФИО не распознано'
    iin=(pension or {}).get('iin') or cd.get('iin') or 'Не распознан'
    date=cd.get('generated_at') or 'Не распознана'
    all_c=list(active)+list(recent)+list(old)+list(revoked)
    load=sum(c.periodic_payment or 0 for c in active)
    out=sum(c.outstanding or 0 for c in active)
    a=Sheet('Анкета')
    a.full('АНКЕТА КЛИЕНТА','title',36)
    a.full(fio,'subtitle',34)
    a.kv('ИИН',str(iin))
    a.kv('Телефон','________________________________')
    a.kv('Дата консультации','________________________________')
    a.section('ЛИЧНЫЕ ДАННЫЕ',150)
    for label in ['Удостоверение / кем выдано','Дата выдачи / срок действия','Адрес проживания','Семейное положение','Супруг(а) / занятость','Дети до 18 лет / иждивенцы']:
        a.kv(label,'________________________________',height=28)
    a.section('ЦЕЛЬ И ФИНАНСИРОВАНИЕ',160)
    for label in ['Цель / программа','Регион / категория / очередь','Дата открытия депозита','Депозит / наличные / БВ / ГП','Оплата: наличные / QR / дата']:
        a.kv(label,'________________________________',height=28)
    a.section('ДАННЫЕ ИЗ ДОКУМЕНТОВ',110)
    a.kv('Средняя зарплата, ₸',srzp if has_pension else 'Выписка не загружена','total')
    a.kv('Платежи по кредитам, ₸',load if has_credit else 'Отчёт не загружен','total')
    a.kv('Дата кредитного отчёта',date if has_credit else 'Отчёт не загружен')
    a.full('Подпись клиента: ____________________    Подпись специалиста: ____________________','note',30)
    a.page()
    a.section('ДОСТУПЫ И ЗАМЕТКИ КОНСУЛЬТАЦИИ',170)
    for label in ['Отбасы Банк: логин / пароль','ЕНПФ: логин / пароль','Кезекте.kz: логин / пароль','ЭЦП','Заметки специалиста']:
        a.kv(label,'________________________________',height=30)
    a.section('ПЕНСИОННЫЕ ОТЧИСЛЕНИЯ',100)
    if not has_pension:
        a.full('Пенсионная выписка не загружена. Доход заполните вручную.','note',38)
    groups=parser.group_pension_by_org((pension or {}).get('rows') or [])
    for key,rs in groups.items():
        name=parser.get_org_display_name(key,(pension or {})['rows'])
        a.section(name,required=min(650,120+len(rs)*22))
        a.kv('БИН / идентификатор',str(key),'note')
        a.row(['Дата поступления',None,'Сумма ОПВ, ₸',None,'Период',None,'Статус'], 'tablehead',26,[(1,2),(3,4),(5,6),(7,8)])
        for rr in rs:
            a.row([rr.date,None,rr.amount,None,str(rr.period or '—'),None,rr.status or '—'],height=22,spans=[(1,2),(3,4),(5,6),(7,8)])
        vals=[float(rr.amount) for rr in rs if rr.amount and rr.amount>0]
        a.kv('СРЗП организации, ₸',parser.calc_pension_avg_vals(vals),'total')
    a.kv('ОБЩАЯ СРЗП, ₸',srzp if has_pension else 'Нет выписки','total',30)
    a.full('Доход учитывается по выбранным организациям. Возвращённые взносы исключены.','note',32)

    c=Sheet('Кредиты');c.full('КРЕДИТНЫЙ ОТЧЁТ','title',36);c.full(fio,'subtitle',34)
    c.kv('Дата актуальности',date if has_credit else 'Отчёт не загружен')
    c.kv('ПКР',cd.get('pkr') if cd.get('pkr') is not None else 'В документе не найден')
    c.kv('Платежи выбранных кредитов, ₸',load if has_credit else 'Нет отчёта','total')
    for title,contracts in [('ДЕЙСТВУЮЩИЕ: УЧТЁННЫЕ КРЕДИТЫ',active),('ЗАВЕРШЁННЫЕ ДО 5 ЛЕТ',recent),('ЗАВЕРШЁННЫЕ БОЛЕЕ 5 ЛЕТ',old),('ОТОЗВАННЫЕ',revoked)]:
        c.section(title,500 if contracts else 80)
        if not contracts:c.full('Договоры отсутствуют в обработанных данных.','note')
        for n,x in enumerate(contracts,1):
            c.section(f'{n}. {x.creditor or "Кредитор не распознан"}',required=430)
            c.kv('Вид финансирования',x.financing_type or 'Не распознан')
            status=x.status or 'Не распознан'
            if re.search(r'Классификация субъекта|Номер договора|Просроченный основной долг',status,re.I):
                status='Не распознан: проверьте в исходном PDF'
            c.kv('Статус договора',status,height=42)
            if parser.contract_has_cession(x):
                c.full('БЫЛА ПЕРЕУСТУПКА ДОЛГА (ЦЕССИЯ)','risk',34)
                c.kv('Цессионарий',x.cessionary or 'Не распознан',height=42)
            for label,value in [('Сумма договора, ₸',x.contract_amount),('Остаток долга, ₸',x.outstanding),('Платёж в месяц, ₸',x.periodic_payment),('Просрочка сейчас: дней',x.current_overdue_days),('Просрочка сейчас: сумма, ₸',x.current_overdue_amount),('Максимальная просрочка: дней',x.max_overdue_days),('Максимальная просрочка: сумма, ₸',x.max_overdue_amount),('Дата завершения',x.end_date)]:
                c.kv(label,value if value is not None else 'Не распознано')

    b=Sheet('Анализ');b.full('АНАЛИЗ ДОКУМЕНТОВ','title',36);b.full(fio,'subtitle',34)
    b.section('ИСТОЧНИКИ И ПОЛНОТА',140)
    b.kv('Кредитный отчёт',('ГКБ / МКБ' if cd.get('source')=='gkb' else 'ПКБ') if has_credit else 'Не загружен')
    b.kv('Дата отчёта',date if has_credit else 'Не загружен')
    b.kv('ИИН кредитного отчёта',str(cd.get('iin') or 'Не распознан'))
    b.kv('ИИН пенсионной выписки',str((pension or {}).get('iin') or 'Не загружена'))
    if has_credit and has_pension and cd.get('iin') and (pension or {}).get('iin') and cd['iin']!=pension['iin']:
        b.full('ИИН источников различаются. Документы обработаны отдельно; объединение не блокируется.','note',42)
    b.section('КЛЮЧЕВЫЕ ПОКАЗАТЕЛИ',220)
    for label,val in [('ПКР',cd.get('pkr') if cd.get('pkr') is not None else 'Не найден'),('Учтено действующих договоров',len(active) if has_credit else 'Нет отчёта'),('Остаток долга, ₸',out if has_credit else 'Нет отчёта'),('Платежи в месяц, ₸',load if has_credit else 'Нет отчёта'),('СРЗП выбранных организаций, ₸',srzp if has_pension else 'Нет выписки'),('Организаций в расчёте',len(groups) if has_pension else 'Нет выписки')]:b.kv(label,val)
    if has_credit and has_pension:
        b.kv('Остаётся после платежей, ₸',srzp-load,'total')
        b.kv('Нагрузка к доходу, %',round(load/srzp*100,1) if srzp>0 else 'Доход равен нулю')
    b.section('РИСК-ПРОФИЛЬ',160)
    now=[x for x in active if (x.current_overdue_days or 0)>0 or (x.current_overdue_amount or 0)>0]
    cess=[x for x in all_c if parser.contract_has_cession(x)]
    b.kv('Текущая просрочка: договоров',len(now) if has_credit else 'Нет отчёта')
    b.kv('Историческая просрочка: дней',max([x.max_overdue_days or 0 for x in all_c]+[0]) if has_credit else 'Нет отчёта')
    worst=max(all_c,key=lambda x:max(x.max_overdue_days or 0,x.current_overdue_days or 0),default=None)
    if worst and max(worst.max_overdue_days or 0,worst.current_overdue_days or 0)>0:
        b.kv('Кредитор с худшей просрочкой',worst.creditor or 'Не распознан',height=42)
    b.kv('Переуступка (цессия): договоров',len(cess) if has_credit else 'Нет отчёта','risk' if cess else 'body')
    missing=sum(x.periodic_payment is None for x in active)
    b.kv('Не распознан платёж: договоров',missing)
    if missing:b.full('Нагрузка может быть неполной. Проверьте платежи указанных договоров.','risk',38)
    b.section('СТРУКТУРА ДОЛГА',110)
    cards=[x for x in active if 'карт' in (x.financing_type or '').lower()]
    b.kv('Кредитные карты: остаток, ₸',sum(x.outstanding or 0 for x in cards))
    b.kv('Остальные займы: остаток, ₸',sum(x.outstanding or 0 for x in active if x not in cards))
    total_amount=sum(x.contract_amount or 0 for x in active)
    if total_amount:b.kv('Погашено от суммы договоров, %',round((total_amount-out)/total_amount*100))
    b.kv('Кредиторов',len({x.creditor for x in active if x.creditor}))
    b.section('ЧТО ДАСТ ДОСРОЧНОЕ ЗАКРЫТИЕ',140)
    limit=None
    if has_credit and has_pension and srzp>0:
        ratio=srzp/4325;kd=.4 if ratio<=40 else .5 if ratio<=65 else .6 if ratio<=90 else .7
        limit=min(round(srzp*kd-load),round(srzp*.9-4325*13-load))
        b.kv('Запас на новый платёж, ₸',limit,'total')
        b.full(f'Справочный сценарий: КД {int(kd*100)}%, прожиточный минимум 56 225 ₸ на 1 человека.','note',38)
    else:b.full('Для сценария нового платежа нужны оба документа и положительный доход.','note',38)
    for x in active:
        if not x.periodic_payment or x.periodic_payment<=0:continue
        b.section(x.creditor or 'Кредитор не распознан',135)
        b.kv('Освободится в месяц, ₸',x.periodic_payment)
        b.kv('Остаток / платёж: месяцев',-(-(x.outstanding or 0)//x.periodic_payment))
        if limit is not None:b.kv('Запас после закрытия, ₸',limit+x.periodic_payment,'total')
    b.full('Срок рассчитан как остаток / платёж без процентов. Это справочный расчёт; решение принимает банк.','note',42)
    return dict(fio=fio,filename=client_filename(fio),sheets=[a.model,c.model,b.model])


def write_printable(output_path,model):
    # Production already uses openpyxl; keep its renderer to avoid adding a JS runtime to Vercel.
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
    from openpyxl.worksheet.page import PageMargins
    from openpyxl.worksheet.pagebreak import Break
    from openpyxl.worksheet.properties import PageSetupProperties
    wb=Workbook();wb.remove(wb.active)
    colors={'section':'E9EDF5','tablehead':'E9EDF5','total':'F0EDF8','risk':'F8E6E8'}
    for spec in model['sheets']:
        ws=wb.create_sheet(spec['name']);ws.sheet_view.showGridLines=False
        for col in 'ABCDEFGH':ws.column_dimensions[col].width=12.2
        for n,row in enumerate(spec['rows'],1):
            ws.row_dimensions[n].height=row['height'];kind=row['kind']
            for j in range(1,9):
                cell=ws.cell(n,j);cell.font=Font(name='Arial',size=18 if kind=='title' else 12 if kind=='subtitle' else 10,bold=kind in ('title','subtitle','section','tablehead','total','risk'),color='202A3D')
                cell.alignment=Alignment(vertical='center',wrap_text=True)
                if kind in colors:cell.fill=PatternFill('solid',fgColor=colors[kind])
                if kind in ('body','total','tablehead'):cell.border=Border(bottom=Side(style='hair',color='CBD1DC'))
                if j<=len(row['values']):
                    val=row['values'][j-1]
                    if isinstance(val,str):cell.number_format='@'
                    cell.value="'"+val if isinstance(val,str) and val.startswith(('=','+','-','@')) else val
                    if isinstance(val,(int,float)):cell.number_format='#,##0.0' if isinstance(val,float) else '#,##0'
        for n,start,end in spec['merges']:ws.merge_cells(start_row=n,start_column=start,end_row=n,end_column=end)
        for r in spec['breaks']:ws.row_breaks.append(Break(id=r))
        ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.orientation='portrait';ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
        ws.sheet_properties.pageSetUpPr=PageSetupProperties(fitToPage=True,autoPageBreaks=False)
        ws.page_margins=PageMargins(left=.35,right=.35,top=.35,bottom=.4,header=.15,footer=.15)
        ws.print_area=f'A1:H{len(spec["rows"])}';ws.print_options.horizontalCentered=True
        ws.oddFooter.center.text='Страница &P из &N';ws.oddFooter.center.size=9
        ws.freeze_panes='A3'
    wb.save(output_path)


def restore_original_questionnaire(output_path, template_path, pension, fio, srzp):
    """Keep the user's questionnaire geometry and fit it to one A4 page."""
    import tempfile
    from copy import copy
    from openpyxl import load_workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.worksheet.page import PageMargins
    from openpyxl.worksheet.properties import PageSetupProperties
    with tempfile.TemporaryDirectory() as directory:
        source=Path(directory)/'questionnaire.xlsx'
        source.write_bytes(Path(template_path).read_bytes())
        if pension:parser.fill_anketa_from_pension(source,pension)
        original=load_workbook(source)['Анкета']
        wb=load_workbook(output_path)
        del wb['Анкета'];ws=wb.create_sheet('Анкета',0)
        for row in original:
            for cell in row:
                dst=ws.cell(cell.row,cell.column,cell.value)
                if cell.has_style:
                    dst.border=copy(cell.border);dst.fill=copy(cell.fill);dst.number_format=cell.number_format
                dst.alignment=Alignment(horizontal=cell.alignment.horizontal,vertical='center',wrap_text=True)
                dst.font=Font(name='Arial',size=9,bold=cell.font.bold,color='202A3D')
        for merged in original.merged_cells:ws.merge_cells(str(merged))
        for col,width in zip('ABCDEFGHIJ',[10,12,10,10,12,10,10,12,10,5]):ws.column_dimensions[col].width=width
        last=max(43,original.max_row if pension else 43)
        for n in range(1,last+1):ws.row_dimensions[n].height=13
        for n in [2,4,10,14,15,17,21,22,23,27,28,44,45]:ws.row_dimensions[n].height=8
        for merged in list(ws.merged_cells):
            if merged.min_row==1:ws.unmerge_cells(str(merged))
        ws.merge_cells('A1:J1');ws['A1']='АНКЕТА КЛИЕНТА';ws['A1'].font=Font(name='Arial',size=15,bold=True);ws['A1'].alignment=Alignment(horizontal='center',vertical='center');ws.row_dimensions[1].height=26
        ws['B5']=fio or 'ФИО не распознано';ws.row_dimensions[5].height=20
        if pension and pension.get('iin'):ws['B6']=str(pension['iin']);ws['B6'].number_format='@'
        ws['A43']='СРЗП';ws['B43']=srzp if pension else 'Выписка не загружена';ws['B43'].number_format='#,##0';ws['B43'].font=Font(name='Arial',size=11,bold=True);ws.row_dimensions[43].height=21
        for band in range(0,max(1,(len(parser.group_pension_by_org((pension or {}).get('rows') or []))+2)//3)):
            top=29 if band==0 else 46+(band-1)*15
            ws.row_dimensions[top].height=42
            if band:
                for left,right in [('A','C'),('D','F'),('G','I')]:
                    rng=f'{left}{top}:{right}{top}'
                    if rng not in str(ws.merged_cells):ws.merge_cells(rng)
            for col in ['A','D','G']:
                ws[f'{col}{top}'].font=Font(name='Arial',size=7,bold=True)
                ws[f'{col}{top}'].alignment=Alignment(wrap_text=True,vertical='center')
            head=top+1
            for offset in [0,3,6]:
                for j,label in enumerate(['Дата','Сумма, ₸','Период']):
                    ws.cell(head,offset+j+1,label);ws.cell(head,offset+j+1).font=Font(name='Arial',size=8,bold=True)
            ws.row_dimensions[head].height=15
        ws.sheet_view.showGridLines=False
        ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.orientation='portrait';ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=1
        ws.sheet_properties.pageSetUpPr=PageSetupProperties(fitToPage=True)
        ws.page_margins=PageMargins(left=.2,right=.2,top=.2,bottom=.2,header=.1,footer=.1)
        ws.print_area=f'A1:J{last}';ws.print_options.horizontalCentered=True;ws.print_options.verticalCentered=False
        wb.save(output_path)
